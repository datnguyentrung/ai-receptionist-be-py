"""Deterministic document parsing, cleanup, validation, and chunking."""

from __future__ import annotations

import hashlib
import io
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path

from app.schemas.ingestion_schema import PreparedChunk

SUPPORTED_SUFFIXES = frozenset({".pdf", ".docx", ".md", ".txt"})
SUPPORTED_MIME_TYPES = {
    ".pdf": {"application/pdf"},
    ".docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
    ".md": {"text/markdown", "text/plain"},
    ".txt": {"text/plain"},
}
CHUNKER_VERSION = "taekwondo-structural-v1"


class DocumentPreprocessingError(ValueError):
    """Raised when deterministic document gates reject an artifact."""


@dataclass(frozen=True)
class SourceBlock:
    text: str
    section: str | None
    page: int | None
    anchor: str


@dataclass(frozen=True)
class PreparedDocument:
    filename: str
    content_hash: str
    normalized_text: str
    chunks: tuple[PreparedChunk, ...]
    warnings: tuple[str, ...] = ()


def prepare_document(
    filename: str,
    data: bytes,
    *,
    max_file_size: int,
    chunk_size_chars: int,
    mime_type: str | None = None,
) -> PreparedDocument:
    if not filename.strip():
        raise DocumentPreprocessingError("Artifact filename is required")
    if not data:
        raise DocumentPreprocessingError("Artifact is empty")
    if len(data) > max_file_size:
        raise DocumentPreprocessingError(
            f"Artifact exceeds the {max_file_size}-byte ingestion limit"
        )

    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise DocumentPreprocessingError(f"Unsupported document type: {suffix or '<none>'}")
    normalized_mime = (mime_type or "").split(";", 1)[0].strip().casefold()
    if normalized_mime and normalized_mime not in SUPPORTED_MIME_TYPES[suffix]:
        raise DocumentPreprocessingError(
            f"MIME type {normalized_mime} does not match the {suffix} artifact"
        )

    blocks = _parse(suffix, filename, data)
    blocks = _deduplicate_blocks(blocks)
    cleaned = tuple(_clean_block(block) for block in blocks if block.text.strip())
    cleaned = tuple(block for block in cleaned if block.text)
    if not cleaned:
        raise DocumentPreprocessingError("Document contains no usable text after cleaning")

    normalized_text = "\n\n".join(block.text for block in cleaned)
    if _valid_character_ratio(normalized_text) < 0.85:
        raise DocumentPreprocessingError("Document text appears corrupted or badly encoded")

    chunks = _chunk(filename, cleaned, max(500, chunk_size_chars))
    if not chunks:
        raise DocumentPreprocessingError("Document could not be divided into chunks")
    return PreparedDocument(
        filename=filename,
        content_hash=_sha256(data),
        normalized_text=normalized_text,
        chunks=tuple(chunks),
    )


def _parse(suffix: str, filename: str, data: bytes) -> tuple[SourceBlock, ...]:
    if suffix in {".md", ".txt"}:
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise DocumentPreprocessingError("Text artifacts must be UTF-8") from exc
        return _parse_structural_text(filename, text)
    if suffix == ".pdf":
        return _parse_pdf(filename, data)
    return _parse_docx(filename, data)


def _parse_structural_text(filename: str, text: str) -> tuple[SourceBlock, ...]:
    section: str | None = None
    blocks: list[SourceBlock] = []
    buffer: list[str] = []

    def flush() -> None:
        if not buffer:
            return
        index = len(blocks)
        blocks.append(
            SourceBlock(
                text="\n".join(buffer),
                section=section,
                page=None,
                anchor=f"{filename}#block-{index}",
            )
        )
        buffer.clear()

    for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        if re.match(r"^#{1,6}\s+\S", line):
            flush()
            section = re.sub(r"^#{1,6}\s+", "", line).strip()
            buffer.append(line)
        elif line.strip():
            buffer.append(line)
        else:
            flush()
    flush()
    return tuple(blocks)


def _parse_pdf(filename: str, data: bytes) -> tuple[SourceBlock, ...]:
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        blocks = []
        for index, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                blocks.append(
                    SourceBlock(
                        text=text,
                        section=None,
                        page=index,
                        anchor=f"{filename}#page-{index}",
                    )
                )
        return tuple(blocks)
    except DocumentPreprocessingError:
        raise
    except Exception as exc:
        raise DocumentPreprocessingError(f"Unable to parse PDF: {exc}") from exc


def _parse_docx(filename: str, data: bytes) -> tuple[SourceBlock, ...]:
    try:
        from docx import Document

        document = Document(io.BytesIO(data))
        blocks: list[SourceBlock] = []
        section: str | None = None
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            if paragraph.style and paragraph.style.name.startswith("Heading"):
                section = text
            blocks.append(
                SourceBlock(
                    text=text,
                    section=section,
                    page=None,
                    anchor=f"{filename}#paragraph-{len(blocks)}",
                )
            )
        for table_index, table in enumerate(document.tables):
            rows = [
                "| " + " | ".join(cell.text.strip() for cell in row.cells) + " |"
                for row in table.rows
            ]
            if rows:
                blocks.append(
                    SourceBlock(
                        text="\n".join(rows),
                        section=section,
                        page=None,
                        anchor=f"{filename}#table-{table_index}",
                    )
                )
        return tuple(blocks)
    except Exception as exc:
        raise DocumentPreprocessingError(f"Unable to parse DOCX: {exc}") from exc


def _deduplicate_blocks(blocks: tuple[SourceBlock, ...]) -> tuple[SourceBlock, ...]:
    seen: set[str] = set()
    result: list[SourceBlock] = []
    for block in blocks:
        normalized = re.sub(r"\s+", " ", unicodedata.normalize("NFKC", block.text)).strip().casefold()
        digest = _sha256(normalized.encode("utf-8"))
        if normalized and digest not in seen:
            seen.add(digest)
            result.append(block)
    return tuple(result)


def _clean_block(block: SourceBlock) -> SourceBlock:
    text = unicodedata.normalize("NFKC", block.text)
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = "\n".join(line.rstrip() for line in text.splitlines())
    text = re.sub(r"(?im)^\s*(?:page|trang)\s+\d+(?:\s+(?:of|/)\s*\d+)?\s*$", "", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return SourceBlock(text=text, section=block.section, page=block.page, anchor=block.anchor)


def _chunk(filename: str, blocks: tuple[SourceBlock, ...], limit: int) -> list[PreparedChunk]:
    groups: list[list[SourceBlock]] = []
    current: list[SourceBlock] = []
    current_size = 0
    for block in blocks:
        if current and current_size + len(block.text) + 2 > limit:
            groups.append(current)
            current = []
            current_size = 0
        if len(block.text) <= limit:
            current.append(block)
            current_size += len(block.text) + 2
            continue
        if current:
            groups.append(current)
            current = []
            current_size = 0
        for start in range(0, len(block.text), limit):
            groups.append(
                [
                    SourceBlock(
                        text=block.text[start : start + limit],
                        section=block.section,
                        page=block.page,
                        anchor=f"{block.anchor}:offset-{start}",
                    )
                ]
            )
    if current:
        groups.append(current)

    chunks: list[PreparedChunk] = []
    for index, group in enumerate(groups):
        text = "\n\n".join(item.text for item in group)
        content_hash = _sha256(text.encode("utf-8"))
        pages = [item.page for item in group if item.page is not None]
        section = next((item.section for item in group if item.section), None)
        chunks.append(
            PreparedChunk(
                chunk_id=_sha256(f"{filename}:{index}:{content_hash}".encode()),
                chunk_index=index,
                text=text,
                content_hash=content_hash,
                token_count=max(1, (len(text) + 3) // 4),
                section=section,
                page_start=min(pages) if pages else None,
                page_end=max(pages) if pages else None,
                source_anchor=",".join(item.anchor for item in group),
            )
        )
    return chunks


def _valid_character_ratio(text: str) -> float:
    if not text:
        return 0.0
    valid = sum(character.isprintable() or character in "\n\t" for character in text)
    return valid / len(text)


def _sha256(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


__all__ = [
    "CHUNKER_VERSION",
    "SUPPORTED_MIME_TYPES",
    "SUPPORTED_SUFFIXES",
    "DocumentPreprocessingError",
    "PreparedDocument",
    "prepare_document",
]
