from io import BytesIO

import pytest
from docx import Document

from app.services.ingestion.preprocessing import (
    DocumentPreprocessingError,
    prepare_document,
)


def prepare(name: str, data: bytes):
    return prepare_document(
        name,
        data,
        max_file_size=1_000_000,
        chunk_size_chars=500,
    )


def test_markdown_chunking_is_deterministic_and_deduplicates_blocks() -> None:
    data = "# Lịch học\n\nThứ Hai 18:00\n\nThứ Hai 18:00\n\n# Học phí\n\n500000".encode()

    first = prepare("lich.md", data)
    second = prepare("lich.md", data)

    assert first.content_hash == second.content_hash
    assert [chunk.chunk_id for chunk in first.chunks] == [
        chunk.chunk_id for chunk in second.chunks
    ]
    assert first.normalized_text.count("Thứ Hai 18:00") == 1
    assert first.chunks[0].source_anchor.startswith("lich.md#block-")


def test_docx_preserves_heading_and_table_anchor() -> None:
    document = Document()
    document.add_heading("Khóa căn bản", level=1)
    document.add_paragraph("Dành cho học viên mới.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Đai"
    table.cell(0, 1).text = "Trắng"
    stream = BytesIO()
    document.save(stream)

    prepared = prepare("khoa-hoc.docx", stream.getvalue())

    assert any(chunk.section == "Khóa căn bản" for chunk in prepared.chunks)
    assert "| Đai | Trắng |" in prepared.normalized_text


@pytest.mark.parametrize("name,data", [("x.csv", b"a,b"), ("x.txt", b"\xff")])
def test_rejects_unsupported_or_invalid_text(name: str, data: bytes) -> None:
    with pytest.raises(DocumentPreprocessingError):
        prepare(name, data)


def test_rejects_mime_type_that_does_not_match_extension() -> None:
    with pytest.raises(DocumentPreprocessingError, match="MIME type"):
        prepare_document(
            "rules.pdf",
            b"not-a-pdf",
            max_file_size=1_000_000,
            chunk_size_chars=500,
            mime_type="text/plain",
        )
